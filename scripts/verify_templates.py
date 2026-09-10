#!/usr/bin/env python3
"""校验 PPT 模板产物是否含有实质内容。

本项目的模板由脚本批量生成，一旦某个生成器没有拿到示例数据，
产出物可能只有标题（甚至一张幻灯片都没有），而生成过程并不会报错。
本脚本把这种「生成成功但内容为空」的情况显式暴露出来，供 CI 拦截。

用法：
    python scripts/verify_templates.py <目录>
    python scripts/verify_templates.py templates --min-shapes 4

判定规则：
    - 零页：幻灯片数为 0
    - 仅标题：幻灯片数 ≥ 1，但首页形状数 < --min-shapes

退出码 0 表示全部通过，1 表示存在不合格产物。
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation

DEFAULT_MIN_SHAPES = 4


def inspect(path: Path) -> tuple[int, int]:
    """返回（幻灯片数，首页形状数）"""
    prs = Presentation(str(path))
    page_count = len(prs.slides)
    shape_count = len(prs.slides[0].shapes) if page_count else 0
    return page_count, shape_count


def main() -> int:
    parser = argparse.ArgumentParser(description="校验 pptx 模板是否含有实质内容")
    parser.add_argument("directory", help="模板目录（会递归查找 *.pptx）")
    parser.add_argument(
        "--min-shapes",
        type=int,
        default=DEFAULT_MIN_SHAPES,
        help=f"首页形状数下限，低于该值判定为仅标题页（默认 {DEFAULT_MIN_SHAPES}）",
    )
    parser.add_argument("-v", "--verbose", action="store_true", help="列出每个不合格文件")
    args = parser.parse_args()

    root = Path(args.directory)
    if not root.is_dir():
        print(f"[FAIL] 目录不存在: {root}", file=sys.stderr)
        return 1

    files = sorted(root.rglob("*.pptx"))
    if not files:
        print(f"[FAIL] 目录中没有找到 .pptx: {root}", file=sys.stderr)
        return 1

    empty: list[str] = []      # 零页
    title_only: list[str] = []  # 只有标题
    groups: Counter[str] = Counter()
    shape_stats: list[int] = []

    for path in files:
        rel = path.relative_to(root).as_posix()
        pages, shapes = inspect(path)
        groups[path.parent.relative_to(root).as_posix()] += 1

        if pages == 0:
            empty.append(rel)
        elif shapes < args.min_shapes:
            title_only.append(f"{rel}（首页 {shapes} 个形状）")
        else:
            shape_stats.append(shapes)

    print(f"扫描目录: {root}")
    print(f"模板总数: {len(files)}    首页形状数下限: {args.min_shapes}")
    print("-" * 60)
    for group, count in sorted(groups.items()):
        print(f"  {group:<28}{count:>6}")
    print("-" * 60)

    failed = len(empty) + len(title_only)
    print(f"零页模板      : {len(empty)}")
    print(f"仅标题模板    : {len(title_only)}")
    if shape_stats:
        print(f"合格模板      : {len(shape_stats)}（首页形状数中位数 {sorted(shape_stats)[len(shape_stats) // 2]}）")

    if args.verbose and (empty or title_only):
        print("\n不合格明细：")
        for rel in empty:
            print(f"  [零页]   {rel}")
        for rel in title_only:
            print(f"  [仅标题] {rel}")

    if failed:
        print(f"\n[FAIL] {failed} 个模板缺少实质内容（共 {len(files)} 个）。", file=sys.stderr)
        return 1

    print(f"\n[OK] {len(files)} 个模板均含实质内容。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
