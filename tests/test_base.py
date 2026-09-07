"""TemplateGenerator 基础类单元测试"""

import os
import sys

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from python.base import TemplateGenerator
from themes import BaseTheme


@pytest.fixture()
def generator():
    return TemplateGenerator(theme=BaseTheme())


def test_default_slide_size(generator):
    assert generator.prs.slide_width == Inches(10)
    assert generator.prs.slide_height == Inches(7.5)


def test_create_slide_increases_slide_count(generator):
    assert len(generator.prs.slides) == 0
    generator.create_slide()
    assert len(generator.prs.slides) == 1
    generator.create_slide()
    assert len(generator.prs.slides) == 2


def test_add_textbox_writes_text_and_font_size(generator):
    generator.create_slide()
    box = generator.add_textbox(
        "Hello PPT", Inches(1), Inches(1), Inches(4), Inches(1), font_size=24
    )
    paragraph = box.text_frame.paragraphs[0]
    assert paragraph.text == "Hello PPT"
    assert paragraph.font.size == Pt(24)


def test_add_textbox_respects_alignment(generator):
    generator.create_slide()
    box = generator.add_textbox(
        "Centered",
        Inches(1),
        Inches(2),
        Inches(4),
        Inches(1),
        alignment=PP_ALIGN.CENTER,
    )
    assert box.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER


def test_add_paragraph_appends_to_existing_frame(generator):
    generator.create_slide()
    box = generator.add_textbox("first", Inches(1), Inches(1), Inches(4), Inches(2))
    generator.add_paragraph(box.text_frame, "second", font_size=14)
    texts = [p.text for p in box.text_frame.paragraphs]
    assert texts == ["first", "second"]


def test_add_shape_applies_solid_fill(generator):
    generator.create_slide()
    shape = generator.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.5),
        Inches(0.5),
        Inches(2),
        Inches(2),
        fill_color=RGBColor(1, 2, 3),
    )
    assert shape.fill.fore_color.rgb == RGBColor(1, 2, 3)


def test_add_gradient_shape_sets_two_stops(generator):
    generator.create_slide()
    shape = generator.add_gradient_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5),
        Inches(0.5),
        Inches(3),
        Inches(2),
        colors=[RGBColor(10, 20, 30), RGBColor(200, 210, 220)],
    )
    stops = shape.fill.gradient_stops
    assert stops[0].color.rgb == RGBColor(10, 20, 30)
    assert stops[1].color.rgb == RGBColor(200, 210, 220)


def test_save_writes_reopenable_pptx(generator, tmp_path):
    generator.create_slide()
    generator.add_textbox("Persisted", Inches(1), Inches(1), Inches(4), Inches(1))
    out = tmp_path / "base.pptx"
    generator.save(str(out))

    assert out.exists()
    assert out.stat().st_size > 0

    reopened = Presentation(str(out))
    assert len(reopened.slides) == 1
    assert reopened.slide_width == Inches(10)


def test_save_returns_filepath(generator, tmp_path):
    generator.create_slide()
    out = tmp_path / "ret.pptx"
    assert generator.save(str(out)) == str(out)
